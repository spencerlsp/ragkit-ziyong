"""DOCX 解析器，基于 python-docx。

注意：第三方包名就叫 ``docx``，和本模块名字很像。
所以本文件特意命名为 docx_parser.py —— 万一你哪天直接
``python src/ragkit/parsing/docx.py`` 当脚本跑，Python 会把文件所在目录
放进 sys.path，``import docx`` 就会 import 到自己，出现莫名其妙的循环导入。
换个名字就彻底没这个坑。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from ..errors import ParseError
from ..schemas import Document
from ..utils import stable_id
from .loaders import run_blocking

__all__ = ["DocxParser"]


def _extract_docx_text(path: Path) -> str:
    """同步抽取 docx 的全部段落文字。"""
    doc = DocxDocument(str(path))

    return "\n\n".join(p.text for p in doc.paragraphs)
    # raise NotImplementedError("TODO: 抽取 docx 段落文字")


class DocxParser:
    """处理 .docx 文件。"""

    extensions: tuple[str, ...] = (".docx",)

    async def parse(self, path: Path) -> Document:
        """解析 .docx 文件。底层异常统一翻译成 ParseError；抽不出文字也报错。"""
        try:
            text = await run_blocking(_extract_docx_text, path)
        except Exception as exc:
            raise ParseError(f"解析 DOCX 失败: {exc}", source=str(path)) from exc

        if not text.strip():
            raise ParseError(
                "DOCX 里没有抽出任何文字",
                source=str(path),
            )

        metadata = {
            "path": str(path),
            "suffix": path.suffix.lower(),
            "size_bytes": path.stat().st_size,
            "char_count": len(text),
        }

        return Document(
            doc_id=stable_id(str(path), text),
            text=text,
            source=str(path),
            metadata=metadata,
        )

        # raise NotImplementedError("TODO: 照 pdf_parser.py 的套路写")
