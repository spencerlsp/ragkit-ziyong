"""解析层的测试。

注意：本项目在 pyproject.toml 里配了 ``asyncio_mode = "auto"``，
所以 ``async def test_xxx`` 会被 pytest-asyncio 自动接管，
不需要写 @pytest.mark.asyncio 装饰器。
"""

from __future__ import annotations

from pathlib import Path

import pytest

from ragkit.errors import ParseError
from ragkit.parsing import parse_file, parse_files, supported_extensions

FIXTURES = Path(__file__).parent / "fixtures"


async def test_parse_text_file(tmp_path: Path) -> None:
    """TODO(你)：
    - 写一个 a.txt，内容 "你好，RAG"
    - doc = await parse_file(path)
    - 断言 doc.text == "你好，RAG"、doc.source == str(path)
    - 断言 doc.metadata["suffix"] == ".txt"
    - 断言 doc.metadata["size_bytes"] > 0
    """
    path = tmp_path / "a.txt"
    path.write_text("你好，RAG", encoding="utf-8")

    doc = await parse_file(path)
    assert doc.text == "你好，RAG"
    assert doc.source == str(path)
    assert doc.metadata["suffix"] == ".txt"
    assert doc.metadata["size_bytes"] > 0


async def test_doc_id_is_deterministic(tmp_path: Path) -> None:
    """TODO(你)：同一个文件解析两次，两次的 doc_id 必须完全相同。
    这条保护的是 M5 的幂等性：doc_id 一旦变成随机的，
    重复导入就会在 Milvus 里堆重复数据，而且很难发现。
    """
    file_path = tmp_path / "stable.txt"
    file_path.write_text("固定内容，用来生成稳定doc_id", encoding="utf-8")

    doc1 = await parse_file(file_path)
    doc2 = await parse_file(file_path)

    assert doc1.doc_id == doc2.doc_id


async def test_unsupported_suffix_raises(tmp_path: Path) -> None:
    """TODO(你)：造一个 .xyz 文件，断言抛 ParseError。
    额外要求：断言异常信息里包含 ".xyz"，并且提到了某个支持的后缀（比如 ".txt"）。
    报错信息本身也是接口的一部分，值得测。
    提示：``with pytest.raises(ParseError) as exc_info:`` 然后 ``str(exc_info.value)``。
    """
    bad_file = tmp_path / "test.xyz"
    bad_file.write_text("随便写点内容", encoding="utf-8")

    with pytest.raises(ParseError) as exc_info:
        await parse_file(bad_file)

    err_msg = str(exc_info.value)
    assert ".xyz" in err_msg
    assert ".txt" in err_msg


async def test_missing_file_raises(tmp_path: Path) -> None:
    """TODO(你)：解析一个不存在的 .txt，断言抛 ParseError 而不是 FileNotFoundError。
    这条在验证「异常在边界被翻译成我们自己的类型」。
    """
    missing_path = tmp_path / "not_exist.txt"
    # 不要创建这个文件

    with pytest.raises(ParseError):
        await parse_file(missing_path)


async def test_supported_extensions_are_sorted() -> None:
    """TODO(你)：断言 supported_extensions() 的结果是排好序的，
    且包含 ".txt" / ".md" / ".pdf" / ".docx"。
    提示：跟 ``tuple(sorted(supported_extensions()))`` 比一下。
    """
    exts = supported_extensions()
    sorted_exts = tuple(sorted(exts))
    assert exts == sorted_exts

    required = {".txt", ".md", ".pdf", ".docx"}
    assert required.issubset(set(exts))


async def test_parse_files_preserves_order(tmp_path: Path) -> None:
    """TODO(你)：传三个**大小差异很大**的文件，断言返回顺序与输入一致。
    造数据建议：
        a = write_text(tmp_path, "a.txt", "A" * 5000)   # 大
        b = write_text(tmp_path, "b.txt", "B")          # 小
        c = write_text(tmp_path, "c.txt", "C")
        docs = await parse_files([a, b, c])
        assert [d.source for d in docs] == [str(a), str(b), str(c)]
    为什么非要「大小差异大」：TaskGroup 不保证完成顺序，小文件几乎肯定先解析完。
    如果实现里用 append 收集结果，这个测试就能抓到 —— 而且是随机翻车的那种失败。
    测试数据要**主动放大实现错误的概率**，否则测试只是走过场。
    """
    a = tmp_path / "a.txt"
    a.write_text("A" * 5000, encoding="utf-8")

    b = tmp_path / "b.txt"
    b.write_text("B", encoding="utf-8")

    c = tmp_path / "c.txt"
    c.write_text("C", encoding="utf-8")

    input_list = [a, b, c]
    docs = await parse_files(input_list)

    sources = [d.source for d in docs]
    expected = [str(p) for p in input_list]
    assert sources == expected


async def test_parse_files_bound_concurrency(tmp_path: Path) -> None:
    """TODO(你)：同一批文件分别用 max_concurrency=1 和 max_concurrency=8 跑，
    两次结果的 doc_id 列表必须完全一样。
    这条测的不是性能（测性能要跑基准，太脆），
    而是「并发度不影响输出内容」—— 顺序和内容都要一致。
    """
    f1 = tmp_path / "f1.txt"
    f1.write_text("文件1", encoding="utf-8")
    f2 = tmp_path / "f2.txt"
    f2.write_text("文件2", encoding="utf-8")
    f3 = tmp_path / "f3.txt"
    f3.write_text("文件3", encoding="utf-8")
    files = [f1, f2, f3]

    docs_serial = await parse_files(files, max_concurrency=1)
    docs_concurrent = await parse_files(files, max_concurrency=8)

    ids_serial = [d.doc_id for d in docs_serial]
    ids_concurrent = [d.doc_id for d in docs_concurrent]
    assert ids_serial == ids_concurrent


async def test_parse_files_empty_input() -> None:
    """TODO(你)：传空列表，应该返回 [] 而不是抛错。
    边界输入是最便宜也最容易漏的测试：空列表、单个元素、重复元素，
    这三类永远值得补一条。
    """
    result = await parse_files([])
    assert result == []


async def test_parse_broken_pdf_raises_parse_error(tmp_path: Path) -> None:
    """TODO(你)：往 .pdf 文件里写一堆垃圾字节，断言抛的是 ParseError。
        bad = tmp_path / "bad.pdf"
        bad.write_bytes(b"this is definitely not a pdf")
    这条测「异常翻译」：pypdf 内部会抛它自己的 PdfReadError 之类，
    但对外必须是 ragkit 的 ParseError，用户 ``except ParseError`` 才兜得住。
    """
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"this is definitely not a pdf")

    with pytest.raises(ParseError):
        await parse_file(bad)


@pytest.mark.skipif(
    not (FIXTURES / "sample.pdf").exists(),
    reason="需要一个真实 PDF：把任意文档「打印成 PDF」另存为 tests/fixtures/sample.pdf",
)
async def test_parse_real_pdf() -> None:
    """TODO(你)：解析 tests/fixtures/sample.pdf，断言：
      - doc.text 非空
      - doc.metadata["page_count"] >= 1
      - doc.metadata["char_count"] == len(doc.text)
    没有这个 PDF 夹具时这条会自动 skip，不影响其他测试。
    """
    pdf_path = FIXTURES / "sample.pdf"
    doc = await parse_file(pdf_path)

    assert doc.text.strip() != ""
    assert doc.metadata["page_count"] >= 1
    assert doc.metadata["char_count"] == len(doc.text)


async def test_parse_docx(tmp_path: Path) -> None:
    """TODO(你)：用 python-docx 现场造一个 docx 当夹具，再解析它。
        from docx import Document as DocxDocument
        d = DocxDocument()
        d.add_paragraph("第一段")
        d.add_paragraph("第二段")
        target = tmp_path / "sample.docx"
        d.save(str(target))
        doc = await parse_file(target)
        assert "第一段" in doc.text
        assert "第二段" in doc.text
        assert "\\n\\n" in doc.text          # 段落之间应该用两个换行分隔
    「在测试里现场造夹具」是个很实用的技巧：不用往仓库塞二进制文件，
    测试还能自解释 —— 看代码就知道输入长什么样。
    """
    from docx import Document as DocxDocument

    d = DocxDocument()
    d.add_paragraph("第一段")
    d.add_paragraph("第二段")
    target = tmp_path / "sample.docx"
    d.save(str(target))

    doc = await parse_file(target)
    assert "第一段" in doc.text
    assert "第二段" in doc.text
    assert "\n\n" in doc.text
